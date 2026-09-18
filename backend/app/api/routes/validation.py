"""
Validation API routes.
Supports: CSV & XLSX upload, per-user job isolation, pause/resume/terminate,
social media column mapping, improved download with per-email statuses.
"""
import csv
import io
import json
import logging
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.models.database import get_db
from app.models.models import ValidationJob, ValidationResult, Template
from app.models.schemas import ValidationJobResponse, ValidationResultResponse
from app.api.routes.auth import get_current_user
from app.models.models import User
from app.worker import process_csv_validation

router = APIRouter()
logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Request models
# ---------------------------------------------------------------------------

class UploadRequest(BaseModel):
    filename: str
    content: str          # Base64 or raw text of the file
    email_column: str = ""
    instagram_column: str = ""
    linkedin_column: str = ""
    whatsapp_column: str = ""
    # For XLSX: sheet name
    sheet_name: str = ""


class SingleValidationRequest(BaseModel):
    email: str


class JobControlRequest(BaseModel):
    action: str  # pause | resume | terminate


# ---------------------------------------------------------------------------
# Single email endpoints
# ---------------------------------------------------------------------------

@router.post("/single")
def validate_single_email(req: SingleValidationRequest):
    from app.core.smtp_socks5 import check_email
    import time
    start = time.time()
    status, reason = check_email(req.email)
    return {
        "email": req.email,
        "status": status,
        "reason": reason,
        "execution_time_ms": round((time.time() - start) * 1000, 2),
    }


@router.post("/deep")
def validate_email_deep(req: SingleValidationRequest):
    """
    Deep email validation with 12-point checks including SMTP verification.
    Returns immediate result (no Celery task).
    """
    from app.core.email_checker import check_email_detailed
    import time
    start = time.time()
    try:
        result = check_email_detailed(req.email)
        result["execution_time_ms"] = round((time.time() - start) * 1000, 2)
        return result
    except Exception as e:
        import traceback
        logger.error(f"Deep validation error for {req.email}: {e}\n{traceback.format_exc()}")
        raise HTTPException(status_code=500, detail=f"Validation failed: {str(e)}")


# ---------------------------------------------------------------------------
# Upload (CSV or XLSX → Celery job)
# ---------------------------------------------------------------------------

@router.post("/upload", response_model=ValidationJobResponse)
async def upload_file(
    req: UploadRequest,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user),
):
    """
    Upload a CSV or XLSX file for bulk validation.
    Dispatches a Celery task and returns the job immediately.
    """
    try:
        filename_lower = req.filename.lower()

        if filename_lower.endswith(".xlsx") or filename_lower.endswith(".xls"):
            # XLSX: decode base64 content and parse with openpyxl
            import base64
            import openpyxl

            raw_bytes = base64.b64decode(req.content)
            wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)

            # Sheet selection
            if req.sheet_name and req.sheet_name in wb.sheetnames:
                ws = wb[req.sheet_name]
            else:
                ws = wb.active

            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                raise HTTPException(status_code=400, detail="XLSX file is empty")

            # First row = headers
            headers = [str(h).strip() if h is not None else f"Col{i}" for i, h in enumerate(rows[0])]
            data_rows = rows[1:]

            # Convert to CSV string for the worker
            csv_buf = io.StringIO()
            writer = csv.DictWriter(csv_buf, fieldnames=headers)
            writer.writeheader()
            for data_row in data_rows:
                row_dict = {}
                for j, h in enumerate(headers):
                    val = data_row[j] if j < len(data_row) else None
                    row_dict[h] = str(val).strip() if val is not None else ""
                writer.writerow(row_dict)
            content_str = csv_buf.getvalue()

        elif filename_lower.endswith(".csv"):
            content_str = req.content

        else:
            raise HTTPException(status_code=400, detail="Only .csv and .xlsx files are supported")

        # Validate email column exists (if provided)
        if req.email_column:
            f = io.StringIO(content_str)
            reader = csv.DictReader(f)
            if reader.fieldnames and req.email_column not in reader.fieldnames:
                raise HTTPException(
                    status_code=400,
                    detail=f"Column '{req.email_column}' not found. Available: {list(reader.fieldnames)}"
                )

        # Count rows for total_records estimate
        f = io.StringIO(content_str)
        reader = csv.DictReader(f)
        total = sum(1 for _ in reader)

        # Build column mapping
        column_mapping = {
            "email": req.email_column,
            "instagram": req.instagram_column,
            "linkedin": req.linkedin_column,
            "whatsapp": req.whatsapp_column,
        }

        # Create job
        job = ValidationJob(
            user_id=current_user.id if current_user else None,
            filename=req.filename,
            total_records=total,
            status="PENDING",
            column_mapping=column_mapping,
        )
        db.add(job)
        db.commit()
        db.refresh(job)

        # Dispatch to Celery
        process_csv_validation.apply_async(
            args=[job.id, content_str, req.email_column],
            kwargs={
                "instagram_column": req.instagram_column,
                "linkedin_column": req.linkedin_column,
                "whatsapp_column": req.whatsapp_column,
            },
            ignore_result=True,
        )

        return job

    except HTTPException:
        raise
    except Exception as e:
        import traceback
        raise HTTPException(status_code=500, detail=f"Internal Error: {str(e)}\n{traceback.format_exc()}")


# ---------------------------------------------------------------------------
# XLSX sheet inspection (called before upload to populate sheet/column selectors)
# ---------------------------------------------------------------------------

class InspectXlsxRequest(BaseModel):
    filename: str
    content: str  # base64


@router.post("/inspect-xlsx")
def inspect_xlsx(req: InspectXlsxRequest):
    """
    Returns sheet names and column names for each sheet of an XLSX file.
    Called client-side to populate sheet/column selectors before uploading.
    """
    try:
        import base64
        import openpyxl
        raw_bytes = base64.b64decode(req.content)
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
        result = {}
        for name in wb.sheetnames:
            ws = wb[name]
            first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            columns = [str(c).strip() if c is not None else f"Col{i}" for i, c in enumerate(first_row)]
            result[name] = columns
        return {"sheets": wb.sheetnames, "columns_by_sheet": result}
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not parse XLSX: {e}")


# ---------------------------------------------------------------------------
# Job listing (filtered by current user if authenticated)
# ---------------------------------------------------------------------------

@router.get("/jobs", response_model=List[ValidationJobResponse])
def get_jobs(
    skip: int = 0,
    limit: int = 200,
    status_filter: Optional[str] = Query(None, alias="status"),
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user),
):
    q = db.query(ValidationJob)
    if current_user:
        q = q.filter(ValidationJob.user_id == current_user.id)
    if status_filter:
        q = q.filter(ValidationJob.status == status_filter)
    jobs = q.order_by(ValidationJob.created_at.desc()).offset(skip).limit(limit).all()
    return jobs


@router.get("/jobs/count")
def get_jobs_count(
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user),
):
    """Return total job count for the current user (dashboard stats)."""
    q = db.query(ValidationJob)
    if current_user:
        q = q.filter(ValidationJob.user_id == current_user.id)
    return {"total": q.count()}


@router.get("/jobs/stats")
def get_jobs_stats(
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user),
):
    """Aggregate stats across ALL jobs for the current user."""
    q = db.query(ValidationJob)
    if current_user:
        q = q.filter(ValidationJob.user_id == current_user.id)
    jobs = q.all()
    total_jobs = len(jobs)
    total_valid = sum(j.valid_count for j in jobs)
    total_invalid = sum(j.invalid_count for j in jobs)
    total_catchall = sum(j.disposable_count for j in jobs)
    total_unknown = sum(j.unknown_count for j in jobs)
    return {
        "total_jobs": total_jobs,
        "total_valid": total_valid,
        "total_invalid": total_invalid,
        "total_catchall": total_catchall,
        "total_unknown": total_unknown,
    }


@router.get("/jobs/{job_id}", response_model=ValidationJobResponse)
def get_job(
    job_id: int,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user),
):
    job = db.query(ValidationJob).filter(ValidationJob.id == job_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    # Skip auth check if user_id is null (allows non-authenticated access)
    if current_user and job.user_id and job.user_id != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")
    return job


# ---------------------------------------------------------------------------
# Job control (pause / resume / terminate)
# ---------------------------------------------------------------------------

@router.post("/jobs/{job_id}/control")
def control_job(
    job_id: int,
    req: JobControlRequest,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user),
):
    job = db.query(ValidationJob).filter(ValidationJob.id == job_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    # Skip auth check if user_id is null (allows non-authenticated access)
    if current_user and job.user_id and job.user_id != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")

    action = req.action.lower()
    if action == "pause":
        if job.status not in ("PROCESSING", "PENDING"):
            raise HTTPException(status_code=400, detail=f"Cannot pause a job in status '{job.status}'")
        job.status = "PAUSED"
    elif action == "resume":
        if job.status != "PAUSED":
            raise HTTPException(status_code=400, detail=f"Cannot resume a job in status '{job.status}'")
        job.status = "PROCESSING"
    elif action == "terminate":
        if job.status in ("COMPLETED", "FAILED", "TERMINATED"):
            raise HTTPException(status_code=400, detail=f"Job is already {job.status}")
        job.status = "TERMINATED"
    else:
        raise HTTPException(status_code=400, detail=f"Unknown action '{action}'. Use: pause, resume, terminate")

    db.commit()
    db.refresh(job)
    return {"job_id": job_id, "status": job.status, "message": f"Job {action}d successfully"}


@router.delete("/jobs/{job_id}")
def delete_job(
    job_id: int,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user),
):
    """Delete a validation job and all its results."""
    job = db.query(ValidationJob).filter(ValidationJob.id == job_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    # Skip auth check if user_id is null (allows non-authenticated access)
    if current_user and job.user_id and job.user_id != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")
    
    db.delete(job)
    db.commit()
    return {"message": "Job deleted successfully", "job_id": job_id}


# ---------------------------------------------------------------------------
# Results
# ---------------------------------------------------------------------------

@router.get("/jobs/{job_id}/results", response_model=List[ValidationResultResponse])
def get_job_results(
    job_id: int,
    skip: int = 0,
    limit: int = 2000,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user),
):
    job = db.query(ValidationJob).filter(ValidationJob.id == job_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    if current_user and job.user_id and job.user_id != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")
    results = (
        db.query(ValidationResult)
        .filter(ValidationResult.job_id == job_id)
        .order_by(ValidationResult.id.asc())
        .offset(skip)
        .limit(limit)
        .all()
    )
    return results


@router.get("/jobs/{job_id}/results/count")
def get_job_results_count(
    job_id: int,
    db: Session = Depends(get_db),
):
    count = db.query(ValidationResult).filter(ValidationResult.job_id == job_id).count()
    return {"count": count}


@router.get("/jobs/{job_id}/columns")
def get_job_columns(
    job_id: int,
    db: Session = Depends(get_db),
):
    first_result = (
        db.query(ValidationResult)
        .filter(ValidationResult.job_id == job_id)
        .first()
    )
    if not first_result or not first_result.original_data:
        raise HTTPException(status_code=404, detail="No results found for this job")
    return {"columns": list(first_result.original_data.keys())}


# ---------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------

@router.get("/jobs/{job_id}/download")
def download_results(
    job_id: int,
    template_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_current_user),
):
    job = db.query(ValidationJob).filter(ValidationJob.id == job_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    if current_user and job.user_id and job.user_id != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")

    results = (
        db.query(ValidationResult)
        .filter(ValidationResult.job_id == job_id)
        .order_by(ValidationResult.id.asc())
        .all()
    )
    if not results:
        raise HTTPException(status_code=400, detail="No results found")

    first = results[0]
    fieldnames = list(first.original_data.keys())
    fieldnames.extend(["Email Status", "Email Check Reason"])

    # Add per-email breakdown column if any row has multiple emails
    has_multi_email = any(r.email_statuses for r in results)
    if has_multi_email:
        fieldnames.append("Per-Email Breakdown")

    # Add social media columns if present
    col_map = job.column_mapping or {}
    if col_map.get("instagram"):
        fieldnames.extend(["Instagram Status", "Instagram Reason"])
    if col_map.get("linkedin"):
        fieldnames.extend(["LinkedIn Status", "LinkedIn Reason"])
    if col_map.get("whatsapp"):
        fieldnames.extend(["WhatsApp Status", "WhatsApp Reason"])

    template = None
    if template_id:
        template = db.query(Template).filter(Template.id == template_id).first()
        if template:
            fieldnames.append("Personalized Content")

    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()

    for result in results:
        row = result.original_data.copy()
        row["Email Status"] = result.status
        row["Email Check Reason"] = result.reason

        if has_multi_email:
            if result.email_statuses:
                breakdown = " | ".join(
                    f"{e['email']}: {e['status']}" for e in result.email_statuses
                )
            else:
                breakdown = f"{result.email}: {result.status}"
            row["Per-Email Breakdown"] = breakdown

        if col_map.get("instagram"):
            row["Instagram Status"] = result.instagram_status or "-"
            row["Instagram Reason"] = result.instagram_reason or "-"
        if col_map.get("linkedin"):
            row["LinkedIn Status"] = result.linkedin_status or "-"
            row["LinkedIn Reason"] = result.linkedin_reason or "-"
        if col_map.get("whatsapp"):
            row["WhatsApp Status"] = result.whatsapp_status or "-"
            row["WhatsApp Reason"] = result.whatsapp_reason or "-"

        if template:
            rendered = template.body_text
            for k, v in result.original_data.items():
                rendered = rendered.replace(f"{{{{{k}}}}}", str(v or ""))
            row["Personalized Content"] = rendered

        writer.writerow(row)

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=validated_{job.filename}"},
    )
